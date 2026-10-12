#!/usr/bin/env python3
"""ja-US: Japanese on a US (101/ANSI) OS keyboard layout.

The same keycaps as ja-JP -- romaji first, the kana as the hint -- but with the
US punctuation, because the OS turns the keys into characters with ITS layout
setting and many Japanese setups run the 101-key one (an English Windows with
Japanese added, a Mac whose keyboard type is ANSI). Letters, digits and
`1 ! 3 # 4 $ 5 % , < . > / ?` are identical; these 14 keys are not.

Writes /tmp/ja_us_cols.json for _patch_xlsx.py. Run it from this directory.
"""
import json, os
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sh = load_workbook(os.path.join(HERE, "lang_lut.xlsx"), data_only=True, read_only=True)["key_lut"]
grid = {(r, c): v for r, row in enumerate(sh.iter_rows(values_only=True), 1)
        for c, v in enumerate(row, 1) if v is not None and v != ""}
langs, i = [], 0
while grid.get((1, 2 + i * 4)):
    langs.append(grid[(1, 2 + i * 4)]); i += 1
col = lambda lang, v: 2 + langs.index(lang) * 4 + v

def cell(v):
    if v is None: return None
    return ["num", int(v)] if isinstance(v, (int, float)) else ["str", str(v)]

ROW = {"KC_2": 29, "KC_6": 33, "KC_7": 34, "KC_8": 35, "KC_9": 36, "KC_0": 37,
       "KC_MINUS": 43, "KC_EQUAL": 44, "KC_LBRC": 45, "KC_RBRC": 46,
       "KC_BACKSLASH": 47, "KC_NONUS_HASH": 48, "KC_SEMICOLON": 49,
       "KC_QUOTE": 50, "KC_GRAVE": 51}

spec = {}
for r in range(2, 71):                       # start from ja-JP, settings included
    spec[str(r)] = [cell(grid.get((r, col("ja-JP", v)))) for v in range(4)]
for kc, r in ROW.items():                    # US base/shift/caps, keep ja-JP's kana hint
    kana = grid.get((r, col("ja-JP", 3)))
    us = [grid.get((r, col("en-US", v))) for v in range(3)]
    spec[str(r)] = [cell(us[0]), cell(us[1]), cell(us[2]), cell(kana)]
spec["48"] = [None, None, None, None]        # no non-US # key on an ANSI board: en-US fallback
spec["51"][3] = None                          # ` ~ is no IME key on US; no kana lives there

json.dump({"ja-US": spec}, open("/tmp/ja_us_cols.json", "w"), ensure_ascii=False, indent=1)
print("wrote /tmp/ja_us_cols.json")
